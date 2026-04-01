from docx import Document

Mem_1 = Document()
f1 = Mem_1.add_paragraph(
    'Кэндибобер, Фейхоаааа, Башорг, Ждун, Фотожаба, Трололо')
Mem_1.save('Mem_1.docx')

Mem_2 = Document()
f2 = Mem_2.add_paragraph(
    'Ждун, Жиза, Кринж, Рофл, Шипперить, Хайп, Пруф, База, Нормис, ЧСВ, Изи,')
Mem_2.save('Mem_2.docx')

def read_set(filepath):
    doc = Document(filepath)
    full_text = []
    for paragraph in doc.paragraphs:
        full_text.append(paragraph.text)
    content = ' '.join(full_text)
    elements = content.replace(',', '').split()
    return set(elements)

def main():
    while True:
        f1 = input("Путь к первому файлу (Mem_1.docx): ")
        A = read_set(f1)
        if A is not None:
            break

    while True:
        f2 = input("Путь ко второму файлу (Mem_2.docx): ")
        B = read_set(f2)
        if B is not None:
            break

    print(f"\nМножество мемов 1: {A}")
    print(f"Множество мемов 2: {B}")

    while True:
        print("\n1 - Объединение")
        print("2 - Пересечение")
        print("3 - Разность")
        print("4 - Симметрическая разность")
        print("0 - Выход")

        choice = input("Выберите операцию: ")

        if choice == '0':
            print("Выход")
            break
        elif choice == '1':
            res = A.union(B)
            print(f"Результат: {res}")
        elif choice == '2':
            res = A.intersection(B)
            print(f"Результат: {res}")
        elif choice == '3':
            res = A.difference(B)
            print(f"Результат: {res}")
        elif choice == '4':
            res = A.symmetric_difference(B)
            print(f"Результат: {res}")
        else:
            print("Неверный выбор")


if __name__ == "__main__":
    main()
